import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.schemas.knowledge_assets import KnowledgeSourceRef

PARSER_VERSION = "oir-excel-parser-v1"
CHUNKING_VERSION = "oir-row-chunking-v1"


@dataclass(frozen=True)
class ParsedKnowledgeChunk:
    chunk_id: str
    content: str
    content_hash: str
    source_ref: KnowledgeSourceRef
    title: str
    structured_payload: dict[str, Any]


@dataclass(frozen=True)
class ParsedWorkbook:
    file_path: Path
    file_hash: str
    stable_key: str
    asset_id: str
    asset_name: str
    chunks: list[ParsedKnowledgeChunk]
    empty: bool
    parser_version: str = PARSER_VERSION
    chunking_version: str = CHUNKING_VERSION


def parse_workbook(path: Path, *, max_chunk_chars: int = 2000) -> ParsedWorkbook:
    data = path.read_bytes()
    file_hash = hashlib.sha256(data).hexdigest()
    stable_key = path.name[:2]
    asset_id = f"asset_content_production_{stable_key}"
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    parsed = []
    for formula_sheet in formulas.worksheets:
        value_sheet = values[formula_sheet.title]
        header_rows = 2 if stable_key == "01" else 1
        headers = _headers(formula_sheet, header_rows)
        for row_number in range(header_rows + 1, formula_sheet.max_row + 1):
            row_values = []
            raw_values = []
            formula_values = {}
            for column in range(1, formula_sheet.max_column + 1):
                formula_value = formula_sheet.cell(row_number, column).value
                display_value = value_sheet.cell(row_number, column).value
                raw_values.append(formula_value)
                if isinstance(formula_value, str) and formula_value.startswith("="):
                    formula_values[get_column_letter(column)] = formula_value
                row_values.append(display_value if display_value is not None else formula_value)
            if not any(_meaningful(value) for value in row_values):
                continue
            fields = {
                headers[index]: _normalize(value)
                for index, value in enumerate(row_values)
                if index < len(headers) and _meaningful(value)
            }
            content = "；".join(f"{key}: {value}" for key, value in fields.items())
            if not content:
                continue
            cell_range = f"A{row_number}:{get_column_letter(formula_sheet.max_column)}{row_number}"
            parts = _split_content(content, max_chunk_chars)
            for part_index, part in enumerate(parts):
                identity = json.dumps(
                    {
                        "file_hash": file_hash,
                        "parser": PARSER_VERSION,
                        "chunking": CHUNKING_VERSION,
                        "sheet": formula_sheet.title,
                        "row": row_number,
                        "part": part_index,
                        "content": part,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                )
                digest = hashlib.sha256(identity.encode()).hexdigest()
                parsed.append(
                    ParsedKnowledgeChunk(
                        chunk_id=f"chunk_content_production_{stable_key}_{digest[:20]}",
                        content=part,
                        content_hash=hashlib.sha256(part.encode()).hexdigest(),
                        title=_title(fields, stable_key, row_number, part_index),
                        source_ref=KnowledgeSourceRef(
                            source_uri=str(path),
                            file_name=path.name,
                            sheet=formula_sheet.title,
                            row_start=row_number,
                            row_end=row_number,
                            cell_range=cell_range,
                            metadata={
                                "columns": list(fields),
                                "formula_cells": formula_values,
                                "raw_values": [_json_value(value) for value in raw_values],
                                "part_index": part_index,
                                "part_count": len(parts),
                                "parser_version": PARSER_VERSION,
                                "chunking_version": CHUNKING_VERSION,
                            },
                        ),
                        structured_payload={
                            "row_id": f"{formula_sheet.title}:{row_number}",
                            "fields": fields,
                            "content": part,
                        },
                    )
                )
    return ParsedWorkbook(
        file_path=path,
        file_hash=file_hash,
        stable_key=stable_key,
        asset_id=asset_id,
        asset_name=path.stem[3:] if len(path.stem) > 3 else path.stem,
        chunks=parsed,
        empty=not parsed,
    )


def _headers(sheet, header_rows: int) -> list[str]:
    headers = []
    for column in range(1, sheet.max_column + 1):
        parts = []
        for row in range(1, header_rows + 1):
            value = _normalize(sheet.cell(row, column).value)
            if value and value not in parts:
                parts.append(value)
        headers.append(" / ".join(parts) or f"column_{column}")
    return headers


def _split_content(content: str, max_chars: int) -> list[str]:
    if len(content) <= max_chars:
        return [content]
    segments = re.split(r"(?<=[。；;\n])", content)
    parts = []
    current = ""
    for segment in segments:
        while len(segment) > max_chars:
            if current:
                parts.append(current)
                current = ""
            parts.append(segment[:max_chars])
            segment = segment[max_chars:]
        if len(current) + len(segment) > max_chars and current:
            parts.append(current)
            current = segment
        else:
            current += segment
    if current:
        parts.append(current)
    return parts


def _title(fields: dict[str, str], stable_key: str, row: int, part: int) -> str:
    for key in ("活动名称", "权益名称", "经营模板", "场景", "要素名称 / 字段名称"):
        if fields.get(key):
            return fields[key]
    return f"{stable_key}-{row}-{part + 1}"


def _meaningful(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return re.sub(r"[ \t]+", " ", text)


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, str | int | float | bool):
        return value
    return str(value)
